from pathlib import Path

import httpx
import openpyxl
import pandas as pd
import pytest
import respx

from notifications.reporter import ExcelReportGenerator
from notifications.telegram_bot import TelegramNotifier, escape_markdown_v2


def test_escape_markdown_v2():
    raw_text = "Preço: €499.99 (Novo!) - Worten [2026]"
    escaped = escape_markdown_v2(raw_text)
    assert r"\." in escaped
    assert r"\(" in escaped
    assert r"\)" in escaped
    assert r"\!" in escaped
    assert r"\-" in escaped
    assert r"\[" in escaped
    assert r"\]" in escaped


@pytest.mark.asyncio
async def test_telegram_notifier_disabled():
    # Quando desativado, retorna True em modo de simulação sem disparar HTTP
    notifier = TelegramNotifier(bot_token="", chat_id="", enabled=False)
    success = await notifier.send_deal_alert(
        product_name="PlayStation 5",
        retailer="worten",
        current_price=399.99,
        avg_price=449.99,
        discount_pct=11.1,
        is_all_time_low=True,
        url="https://worten.pt/ps5",
    )
    assert success is True


@pytest.mark.asyncio
@respx.mock
async def test_telegram_notifier_active_send():
    token = "test_bot_token"
    chat_id = "12345678"
    respx.post(f"https://api.telegram.org/bot{token}/sendMessage").mock(
        return_value=httpx.Response(200, json={"ok": True})
    )

    notifier = TelegramNotifier(bot_token=token, chat_id=chat_id, enabled=True)
    success = await notifier.send_deal_alert(
        product_name="MacBook Air M3",
        retailer="amazon",
        current_price=1099.00,
        avg_price=1299.00,
        discount_pct=15.4,
        is_all_time_low=True,
        url="https://amazon.es/dp/123",
    )
    assert success is True


def test_excel_report_generation(tmp_path: Path):
    df = pd.DataFrame(
        [
            {
                "product_id": 1,
                "retailer": "worten",
                "sku": "SKU-EXCEL-1",
                "name": "Nintendo Switch OLED",
                "price": 299.99,
                "historical_avg": 349.99,
                "historical_min": 299.99,
                "discount_vs_avg_pct": 14.3,
                "is_all_time_low": True,
                "is_deal_alert": True,
                "availability": "in_stock",
                "url": "https://worten.pt/switch",
            }
        ]
    )

    generator = ExcelReportGenerator(output_dir=tmp_path)
    file_path = generator.generate_intelligence_report(df, filename_prefix="test_report")

    assert file_path.exists()
    assert file_path.suffix == ".xlsx"

    # Abre o ficheiro gerado e valida estrutura das células
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    assert ws is not None
    assert "APEX PRICE ENGINE" in str(ws["A1"].value)
    assert ws["D5"].value == "Nintendo Switch OLED"
    assert ws["E5"].value == 299.99
    assert ws["I5"].value == "SIM"
