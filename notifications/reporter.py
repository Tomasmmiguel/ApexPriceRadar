import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


def _safe_float(value: Any, default: float = 0.0) -> float:
    """Converte valores numéricos com segurança evitando erros com None ou NaN."""
    if value is None or pd.isna(value):
        return default
    try:
        return float(value)
    except (ValueError, TypeError):
        return default


class ExcelReportGenerator:
    """
    Gera relatórios executivos em formato Microsoft Excel (.xlsx) estilizados
    com identidade corporativa, formatação condicional e tipos nativos de moeda.
    """

    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate_intelligence_report(
        self, df: pd.DataFrame, filename_prefix: str = "relatorio_precos"
    ) -> Path:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_path = self.output_dir / f"{filename_prefix}_{timestamp}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        if not isinstance(ws, Worksheet):
            ws = wb.create_sheet()

        ws.title = "Preços & Inteligência"
        ws.views.sheetView[0].showGridLines = True

        # Paleta de Estilos Corporativos
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        deal_fill = PatternFill(
            start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
        )  # Verde suave
        out_stock_fill = PatternFill(
            start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
        )  # Vermelho suave

        regular_font = Font(name="Segoe UI", size=10)
        bold_font = Font(name="Segoe UI", size=10, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # 1. Título do Relatório
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "APEX PRICE ENGINE - RELATÓRIO EXECUTIVO DE INTELIGÊNCIA DE PREÇOS"
        title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 28

        # 2. Metadados de Geração
        ws["A2"] = (
            f"Extraído em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Total de Produtos: {len(df)}"
        )
        ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="595959")
        ws.row_dimensions[2].height = 18

        # 3. Cabeçalhos das Colunas
        headers = [
            "ID",
            "Loja",
            "SKU",
            "Nome do Produto",
            "Preço Atual (€)",
            "Média Histórica (€)",
            "Mín. Histórico (€)",
            "Desconto vs Média",
            "Mínimo Absoluto?",
            "Estado",
            "Link",
        ]
        ws.append([])  # Linha 3 vazia

        ws.append(headers)  # Linha 4
        ws.row_dimensions[4].height = 24

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # 4. Dados
        if not df.empty:
            for _, row in df.iterrows():
                is_deal = bool(row.get("is_deal_alert", False))
                availability = str(row.get("availability", "in_stock"))
                is_all_time_low = bool(row.get("is_all_time_low", False))

                current_p = _safe_float(row.get("price"))
                avg_p = _safe_float(row.get("historical_avg"), default=current_p)
                min_p = _safe_float(row.get("historical_min"), default=current_p)
                discount_val = _safe_float(row.get("discount_vs_avg_pct")) / 100.0

                row_values = [
                    int(row.get("product_id", 0)),
                    str(row.get("retailer", "")).upper(),
                    str(row.get("sku", "")),
                    str(row.get("name", "")),
                    current_p,
                    avg_p,
                    min_p,
                    discount_val,
                    "SIM" if is_all_time_low else "NÃO",
                    "Disponível" if availability == "in_stock" else "Esgotado",
                    str(row.get("url", "")),
                ]
                ws.append(row_values)
                current_row = ws.max_row
                ws.row_dimensions[current_row].height = 20

                # Formatação e cores das células
                for col_idx in range(1, len(row_values) + 1):
                    c = ws.cell(row=current_row, column=col_idx)
                    c.font = regular_font
                    c.border = thin_border

                    # Formatos numéricos específicos
                    if col_idx in (5, 6, 7):  # Preços
                        c.number_format = '"€" #,##0.00'
                        c.alignment = Alignment(horizontal="right")
                    elif col_idx == 8:  # Desconto %
                        c.number_format = "0.0%"
                        c.alignment = Alignment(horizontal="right")
                        if row.get("discount_vs_avg_pct", 0.0) > 0:
                            c.font = bold_font
                    elif col_idx in (1, 2, 3, 9, 10):
                        c.alignment = Alignment(horizontal="center")

                    # Destaques de linha
                    if is_deal:
                        c.fill = deal_fill
                    elif availability != "in_stock":
                        c.fill = out_stock_fill

        # 5. Ajuste Automático de Largura de Coluna
        for col in ws.columns:
            first_cell = col[0]
            col_idx = first_cell.column
            if not isinstance(col_idx, int):
                continue
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for cell in col:
                if cell.row in (1, 2, 3):
                    continue
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 60)

        wb.save(file_path)
        logger.info(f"Relatório Excel executivo gerado com sucesso em: {file_path}")
        return file_path
